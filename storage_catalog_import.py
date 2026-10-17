"""Strict, read-only importer for the OPS A:I storage worksheet."""
import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def name_key(value):
    return clean(value).casefold().replace('ё', 'е')


def hours(value):
    text = clean(value).casefold().replace(',', '.')
    if not text or text == '-' or 'согласно' in text:
        return None
    # Parentheses describe temperature, tempering or container, not extra shelf life.
    text = text.split('(')[0]
    if re.search(r'[-−]\s*\d', text):
        raise ValueError(f'Отрицательный или интервальный срок: {value}')
    if re.search(r'\d\s*[-–/]\s*\d', text):
        raise ValueError(f'Неоднозначный срок: {value}')
    matches = re.findall(r'(\d+(?:\.\d+)?)\s*(день|дня|дней|дей|сутки|суток|сут|часа|часов|час|ч\.?|минуты|минут|минута|мин\.?)', text)
    if not matches:
        raise ValueError(f'Не удалось прочитать срок: {value}')
    total = sum(float(number) * (24 if unit.startswith(('д', 'сут')) else 1/60 if unit.startswith('мин') else 1) for number, unit in matches)
    if not 0 < total <= 24 * 365:
        raise ValueError(f'Недопустимый срок: {value}')
    return total


def production_variants(value):
    text = clean(value)
    if not text or text == '-' or 'согласно' in text.casefold():
        return []
    pieces = [piece.strip(' ;') for piece in re.split(r'(?<![\w/])(?=(?:ТД|Д)\s*=)', text) if piece.strip(' ;')]
    return [{'label': piece, 'hours': hours(piece)} for piece in pieces]


def parse_storage(path):
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or path.stat().st_size > 10_000_000:
        raise ValueError('Нужен файл .xlsx размером до 10 МБ.')
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 50_000_000:
            raise ValueError('Слишком большой распакованный Excel.')
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook['Общее'] if 'Общее' in workbook.sheetnames else workbook.active
        if sheet.max_row > 5000:
            raise ValueError('В таблице больше 5000 строк.')
        header = [clean(v).casefold() for v in next(sheet.iter_rows(min_row=5, max_row=5, max_col=9, values_only=True))]
        if header[1] != 'общий срок хранения' or header[2] != 'размораживание' or header[7] != 'маркировка':
            raise ValueError('Структура не совпадает с шаблоном OPS: ожидаются колонки A–I и заголовки в строке 5.')
        rows, errors, warnings, seen = [], [], [], set()
        for number, cells in enumerate(sheet.iter_rows(min_row=6, max_col=9), 6):
            values = [clean(cell.value) for cell in cells]
            name, prep, kind, defrost, storage, label, place, production, production_place = values
            if not name:
                continue
            if name.startswith(('П =', 'T =', 'Т =', 'Д =')):
                break
            if production.casefold() == 'маркировка':
                continue  # Section header, not a product.
            try:
                if any(len(value) > 3000 for value in values):
                    raise ValueError('Слишком длинное значение в ячейке.')
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError('Формулы не поддерживаются. Пришлите таблицу со значениями.')
                if name_key(name) in seen:
                    raise ValueError(f'Повторяется продукт: {name}')
                seen.add(name_key(name))
                frozen = kind in {'*', '**'}
                defrost_hours = hours(defrost) if frozen else None
                storage_hours = hours(storage) if frozen else None
                prep_hours = hours(prep)
                if frozen and prep_hours is None and defrost_hours and storage_hours:
                    prep_hours = defrost_hours + storage_hours
                if frozen and all(v is not None for v in (prep_hours, defrost_hours, storage_hours)) and abs(prep_hours - defrost_hours - storage_hours) > .001:
                    raise ValueError('Общий срок не равен размораживанию + хранению.')
                if not frozen:
                    hours(label)  # Validate the opened-package duration too.
                variants = production_variants(production)
                if not any([prep_hours, hours(label) if not frozen else None, variants]):
                    if 'согласно' not in production.casefold():
                        raise ValueError('Не указано ни одного распознаваемого срока.')
                    warnings.append(f'{name}: срок только по упаковке, без автоматического расчёта.')
                if 'дей' in label.casefold():
                    warnings.append(f'{name}: «дей» прочитано как «дней».')
                row = dict(name=name, prepHours=prep_hours, defrostType=kind if frozen else '',
                           openedHours=hours(label) if not frozen else None,
                           defrostHours=defrost_hours, storageHours=storage_hours,
                           prepRaw=prep, defrostRaw=defrost if frozen else '', storageRaw=storage,
                           prepLabel=label, prepPlace=place, productionHours=variants[0]['hours'] if len(variants) == 1 else None,
                           productionLabel=production, productionPlace=production_place, sourceRow=number,
                           productionVariants=variants)
                rows.append(row)
            except ValueError as error:
                errors.append(f'Строка {number} ({name}): {error}')
        if errors:
            raise ValueError('\n'.join(errors[:15]))
        if not rows:
            raise ValueError('В файле не найдены продукты.')
        if len(rows) > 2000:
            raise ValueError('В каталоге допустимо не более 2000 продуктов.')
        return {'items': rows, 'warnings': warnings, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    finally:
        workbook.close()


def baseline():
    text = Path('docs/storage-times.js').read_text(encoding='utf-8')
    return json.loads(text.split('=', 1)[1].strip().rstrip(';'))


def merge_catalog(existing, imported):
    result = [dict(row) for row in existing]
    positions = {name_key(row['name']): i for i, row in enumerate(result)}
    changes = []
    fields = ('prepRaw', 'defrostRaw', 'storageRaw', 'defrostType', 'prepLabel', 'prepPlace', 'productionLabel', 'productionPlace',
              'prepHours', 'defrostHours', 'storageHours', 'productionHours')
    for row in imported:
        index = positions.get(name_key(row['name']))
        old = result[index] if index is not None else None
        changed = [field for field in fields if clean((old or {}).get(field)) != clean(row.get(field))]
        if old is None or changed:
            changes.append({'name': row['name'], 'old': old, 'new': row, 'fields': changed})
        if index is None:
            positions[name_key(row['name'])] = len(result)
            result.append(row)
        else:
            result[index] = row
    retained = sum(name_key(row['name']) not in {name_key(item['name']) for item in imported} for row in existing)
    return result, changes, retained
